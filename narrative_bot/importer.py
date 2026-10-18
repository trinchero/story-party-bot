"""The workbook is data, never executable instructions."""
from io import BytesIO
from itertools import islice
import re
from zipfile import ZipFile

from openpyxl import load_workbook

HEADERS = {
    'Guia': ['Campo', 'Contenido'],
    'Personajes': ['Código', 'Nombre', 'Descripción', 'Ficha privada'],
    'Fases': ['Fase', 'Título', 'Notas admin'],
    'Mensajes': ['Fase', 'Orden', 'Destinatarios', 'Título admin', 'Texto', 'Notas admin'],
}


class InvalidWorkbook(ValueError):
    pass


def parse_excel(data: bytes) -> dict:
    errors = []

    def error(sheet, row, correction):
        errors.append(f'{sheet}, fila {row}: {correction}')

    if len(data) > 5_000_000:
        raise InvalidWorkbook('Archivo, fila 1: usa un Excel de menos de 5 MB.')
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 30_000_000:
                raise InvalidWorkbook('Archivo, fila 1: reduce el tamaño descomprimido a 30 MB.')
        wb = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except InvalidWorkbook:
        raise
    except Exception as exc:
        raise InvalidWorkbook('Archivo, fila 1: adjunta un archivo .xlsx válido.') from exc
    tables = {}
    try:
        for name, headers in HEADERS.items():
            if name not in wb.sheetnames:
                error(name, 1, 'añade esta hoja con sus encabezados.')
                continue
            sheet = wb[name]
            if (sheet.max_row or 0) > 3000 or (sheet.max_column or 0) > 30:
                error(name, 1, 'reduce la hoja a 3000 filas y 30 columnas.')
                continue
            # Some exporters omit or understate worksheet dimensions.
            sheet.reset_dimensions()
            rows = list(islice(sheet.iter_rows(max_col=31), 3001))
            if len(rows) > 3000 or any(row[30].value is not None for row in rows):
                error(name, 1, 'reduce la hoja a 3000 filas y 30 columnas.')
                continue
            actual = [str(c.value or '').strip() for c in rows[0]] if rows else []
            if any(actual.count(h) != 1 for h in headers):
                error(name, 1, 'incluye una sola vez las columnas: ' + ', '.join(headers))
                continue
            records = []
            for number, cells in enumerate(rows[1:], 2):
                if all(c.value is None for c in cells):
                    continue
                record = {'row': number}
                for h in headers:
                    c = cells[actual.index(h)]
                    if c.data_type == 'f':
                        error(name, number, f'sustituye la fórmula en {h} por texto o un número.')
                    record[h] = str(c.value).strip() if c.value is not None else ''
                records.append(record)
            tables[name] = records
    except Exception as exc:
        raise InvalidWorkbook('Archivo, fila 1: vuelve a guardar el archivo como .xlsx; no se pudieron leer sus hojas.') from exc
    finally:
        wb.close()
    if errors:
        raise InvalidWorkbook('\n'.join(errors))
    return validate_tables(tables)


def validate_tables(tables, headers=None):
    headers = headers or HEADERS
    errors = []

    def error(sheet, row, correction):
        errors.append(f'{sheet}, fila {row}: {correction}')

    for sheet, records in tables.items():
        if not records:
            error(sheet, 2, 'añade al menos una fila.')
        for r in records:
            for h in headers[sheet]:
                if h != 'Notas admin' and not r[h]:
                    error(sheet, r['row'], f'completa {h}.')
            for h in ('Texto', 'Ficha privada'):
                if h in r and len(r[h].encode('utf-16-le')) // 2 > 4000:
                    error(sheet, r['row'], f'reduce {h} a 4000 caracteres Telegram.')
    codes = set()
    for r in tables['Personajes']:
        code = r['Código']
        if not re.fullmatch(r'[A-Z0-9_-]{1,32}', code) or code == 'TODOS':
            error('Personajes', r['row'], 'usa un código de 1–32 letras MAYÚSCULAS, cifras, _ o -; distinto de TODOS.')
        if code in codes:
            error('Personajes', r['row'], 'usa un código único.')
        codes.add(code)

    def integer(r, key, sheet):
        try:
            n = int(r[key])
            if n < 1:
                raise ValueError
            r[key] = n
        except ValueError:
            error(sheet, r['row'], f'{key} debe ser un entero positivo.')
            r[key] = 0

    phases = []
    for r in tables['Fases']:
        integer(r, 'Fase', 'Fases')
        phases.append(r['Fase'])
    if sorted(phases) != list(range(1, len(phases) + 1)):
        for r in tables['Fases']:
            error('Fases', r['row'], 'numera las fases 1, 2, 3… sin saltos ni duplicados.')
    for r in tables['Mensajes']:
        integer(r, 'Fase', 'Mensajes')
        integer(r, 'Orden', 'Mensajes')
        if r['Fase'] not in phases:
            error('Mensajes', r['row'], 'indica una fase existente en Fases.')
        recipients = [x.strip() for x in r['Destinatarios'].split(',')]
        if recipients != ['TODOS'] and (not all(x in codes for x in recipients) or len(set(recipients)) != len(recipients)):
            error('Mensajes', r['row'], 'usa TODOS solo, o códigos existentes separados por comas, sin repetir.')
        r['recipients'] = recipients
    for phase in phases:
        messages = [r for r in tables['Mensajes'] if r['Fase'] == phase]
        if not messages:
            error('Mensajes', 2, f'añade al menos un mensaje para la fase {phase}.')
        if sorted(r['Orden'] for r in messages) != list(range(1, len(messages) + 1)):
            for r in messages:
                error('Mensajes', r['row'], f'numera Orden en fase {phase} desde 1 sin saltos ni duplicados.')
    if errors:
        raise InvalidWorkbook('\n'.join(errors))
    tables['Fases'].sort(key=lambda r: r['Fase'])
    tables['Mensajes'].sort(key=lambda r: (r['Fase'], r['Orden']))
    return tables


def parse_numbers(data: bytes) -> dict:
    """Read native Numbers tables directly; no Excel conversion or application needed."""
    from numbers_parser import Document
    from tempfile import TemporaryDirectory
    from pathlib import Path
    if len(data) > 5_000_000:
        raise InvalidWorkbook('Archivo, fila 1: usa un archivo Numbers de menos de 5 MB.')
    required = {
        'Personajes': ['Código', 'Nombre'],
        'Fases': ['Fase', 'Título'],
        'Mensajes': ['Fase', 'Orden', 'Destinatarios', 'Texto'],
    }
    errors, tables = [], {}
    def error(sheet, row, correction):
        errors.append(f'{sheet}, fila {row}: {correction}')
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 30_000_000:
                raise InvalidWorkbook('Archivo, fila 1: reduce el tamaño descomprimido a 30 MB.')
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'story.numbers'
            path.write_bytes(data)
            document = Document(path)
            for name, headers in required.items():
                matches = [s for s in document.sheets if s.name == name]
                if len(matches) != 1:
                    error(name, 1, 'añade una sola hoja con este nombre.')
                    continue
                sheet = matches[0]
                if len(sheet.tables) != 1:
                    error(name, 1, 'usa una sola tabla en esta hoja.')
                    continue
                table = sheet.tables[0]
                if table.num_rows > 3000 or table.num_cols > 30:
                    error(name, 1, 'reduce la tabla a 3000 filas y 30 columnas.')
                    continue
                rows = table.rows()
                actual = [str(c.value or '').strip() for c in rows[0]] if rows else []
                if any(actual.count(h) != 1 for h in headers):
                    error(name, 1, 'incluye una sola vez las columnas: ' + ', '.join(headers))
                    continue
                records = []
                for number, cells in enumerate(rows[1:], 2):
                    if all(c.value in (None, '') and not c.is_formula for c in cells):
                        continue
                    record = {'row': number}
                    for h in headers:
                        cell = cells[actual.index(h)]
                        if cell.is_formula:
                            error(name, number, f'sustituye la fórmula en {h} por texto o un número.')
                        value = cell.value
                        if isinstance(value, float) and value.is_integer():
                            value = int(value)
                        record[h] = str(value).strip() if value is not None else ''
                    records.append(record)
                tables[name] = records
    except InvalidWorkbook:
        raise
    except Exception as exc:
        raise InvalidWorkbook('Archivo, fila 1: guarda un archivo .numbers válido, sin contraseña y como archivo único (no paquete).') from exc
    if errors:
        raise InvalidWorkbook('\n'.join(errors))
    tables = validate_tables(tables, required)
    # Keep the internal story shape compatible with previously stored games.
    tables['Guia'] = []
    for row in tables['Fases']:
        row['Notas admin'] = ''
    for row in tables['Mensajes']:
        row['Título admin'] = row['Texto'].splitlines()[0][:70]
        row['Notas admin'] = ''
    return tables
